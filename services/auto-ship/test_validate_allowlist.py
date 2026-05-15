"""Tests for services/auto-ship/validate_allowlist.py.

Each rule is exercised positively (clean case passes) and negatively
(violating case fails with the expected rule name). Uses pytest tmp_path
to build synthetic skill trees so the tests are hermetic.
"""

from __future__ import annotations

import textwrap
from pathlib import Path

import pytest

import validate_allowlist as v


# ---------------------------------------------------------------------------
# Helpers — build synthetic skill trees in tmp_path
# ---------------------------------------------------------------------------

def _write_skill(
    skills_root: Path,
    name: str,
    *,
    description: str = "A clean Tier 1 skill.",
    python_files: dict[str, str] | None = None,
) -> Path:
    """Create skills_root/<name>/SKILL.md + optional py files. Returns the dir."""
    skill_dir = skills_root / name
    skill_dir.mkdir(parents=True, exist_ok=True)
    skill_md = skill_dir / "SKILL.md"
    skill_md.write_text(
        textwrap.dedent(
            f"""\
            ---
            name: {name}
            description: {description}
            ---

            # {name}
            """
        ),
        encoding="utf-8",
    )
    if python_files:
        for relpath, content in python_files.items():
            target = skill_dir / relpath
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(content, encoding="utf-8")
    return skill_dir


def _write_manifest(path: Path, **sections: list[str]) -> Path:
    import yaml

    path.write_text(yaml.safe_dump(sections), encoding="utf-8")
    return path


@pytest.fixture
def automations(tmp_path: Path) -> Path:
    """A fake blackstone-automations checkout with a skills/ subdir."""
    skills = tmp_path / "automations" / "skills"
    skills.mkdir(parents=True)
    return skills.parent


# ---------------------------------------------------------------------------
# Happy path
# ---------------------------------------------------------------------------

def test_clean_manifest_passes(tmp_path: Path, automations: Path):
    skills = automations / "skills"
    _write_skill(skills, "discovery-tracker-sync")
    _write_skill(skills, "alarm-system")
    _write_skill(skills, "mc-letter")

    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["discovery-tracker-sync", "alarm-system"],
        excluded=["mc-letter"],
    )
    assert v.validate(manifest, skills) == []


# ---------------------------------------------------------------------------
# Rule: exclusive — entry must be in exactly one section
# ---------------------------------------------------------------------------

def test_entry_in_two_sections_fails(tmp_path: Path, automations: Path):
    skills = automations / "skills"
    _write_skill(skills, "foo")
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["foo"],
        probe_only=["foo"],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "exclusive" and x.skill == "foo" for x in vs)


# ---------------------------------------------------------------------------
# Rule: exists — skill must exist in blackstone-automations
# ---------------------------------------------------------------------------

def test_missing_skill_fails(tmp_path: Path, automations: Path):
    skills = automations / "skills"
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["does-not-exist"],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "exists" and x.skill == "does-not-exist" for x in vs)


# ---------------------------------------------------------------------------
# Rule: drafter-name — name regex catches obvious drafters
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "name",
    ["mc-letter", "mediation-brief", "declaration-iso-mtc", "rfp-drafting",
     "separate-statement-mtc", "rfp-responses", "legal-doc-drafting",
     "stip-to-stay", "proof-of-service"],
)
def test_drafter_name_blocks_in_auto_loop(tmp_path: Path, automations: Path, name: str):
    skills = automations / "skills"
    _write_skill(skills, name)
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=[name],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "drafter-name" and x.skill == name for x in vs)


def test_drafter_name_does_not_block_in_excluded(tmp_path: Path, automations: Path):
    """The drafter-name rule applies only to auto_loop + probe_only entries."""
    skills = automations / "skills"
    _write_skill(skills, "mc-letter")
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        excluded=["mc-letter"],
    )
    vs = v.validate(manifest, skills)
    assert not any(x.rule == "drafter-name" for x in vs)


# ---------------------------------------------------------------------------
# Rule: drafter-description — SKILL.md frontmatter description
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "phrase",
    [
        "Drafts a letter to opposing counsel.",
        "Drafting RFP responses overnight.",
        "Auto-draft hook generates email replies.",
        "Produces drafts for attorney review.",
        "Generates a draft mediation brief.",
        "Composes a pleading for filing.",
        "Adds the firm letterhead automatically.",
        "Standalone motion (filing) submitter.",
    ],
)
def test_drafter_description_fails(tmp_path: Path, automations: Path, phrase: str):
    skills = automations / "skills"
    _write_skill(skills, "trojan", description=phrase)
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["trojan"],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "drafter-description" and x.skill == "trojan" for x in vs)


@pytest.mark.parametrize(
    "phrase",
    [
        # Real Tier 1 idioms that previously false-positived on bare "draft".
        "Routes queued discovery-response drafts to the per-doc-type drafter skills.",
        "Watches for pending follow-up drafts and surfaces them for approval.",
        "Tracks motion outcomes and pleadings filed by opposing counsel.",
    ],
)
def test_drafter_description_tolerates_orchestration_language(
    tmp_path: Path, automations: Path, phrase: str
):
    """Bare 'drafts' (noun, referring to other skills' output) must not fire."""
    skills = automations / "skills"
    _write_skill(skills, "orchestrator", description=phrase)
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["orchestrator"],
    )
    vs = v.validate(manifest, skills)
    assert not any(x.rule == "drafter-description" for x in vs), vs


def test_drafter_content_rules_skip_probe_only(tmp_path: Path, automations: Path):
    """probe_only entries are watched, not touched — drafter content rules
    (description, python-docx, output-path) do not apply to them. Only the
    cheap name-regex still fires for safety."""
    skills = automations / "skills"
    _write_skill(
        skills,
        "watcher-with-draft-hook",
        description="Monitors task runs and auto-drafts a status email.",
        python_files={
            "build.py": (
                "from docx import Document\n"
                "def build():\n"
                "    Document().save('out.docx')\n"
            ),
        },
    )
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        probe_only=["watcher-with-draft-hook"],
    )
    vs = v.validate(manifest, skills)
    # Name-regex doesn't match this name, so the only checks that could
    # fire are the content rules — and they're scoped out for probe_only.
    assert vs == [], vs


# ---------------------------------------------------------------------------
# Rule: python-docx-output — AST scan for docx import + Document()
# ---------------------------------------------------------------------------

def test_python_docx_output_fails(tmp_path: Path, automations: Path):
    skills = automations / "skills"
    _write_skill(
        skills,
        "trojan-drafter",
        python_files={
            "build.py": textwrap.dedent(
                """
                from docx import Document
                def build():
                    doc = Document()
                    doc.add_paragraph("hi")
                    doc.save("out.docx")
                """
            )
        },
    )
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["trojan-drafter"],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "python-docx-output" and x.skill == "trojan-drafter" for x in vs)


def test_python_docx_not_imported_passes(tmp_path: Path, automations: Path):
    skills = automations / "skills"
    _write_skill(
        skills,
        "clean-syncer",
        python_files={
            "sync.py": "def run():\n    return 1\n",
        },
    )
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["clean-syncer"],
    )
    assert v.validate(manifest, skills) == []


# ---------------------------------------------------------------------------
# Rule: drafter-output-path — string-literal scan
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("fragment", ["Active Cases/", "/Drafts/", "/Pleadings/"])
def test_drafter_output_path_fails(tmp_path: Path, automations: Path, fragment: str):
    skills = automations / "skills"
    _write_skill(
        skills,
        "sneaky",
        python_files={
            "write.py": f'PATH = "{fragment}thing"\n',
        },
    )
    manifest = _write_manifest(
        tmp_path / "skills_eligible.yaml",
        auto_loop=["sneaky"],
    )
    vs = v.validate(manifest, skills)
    assert any(x.rule == "drafter-output-path" and x.skill == "sneaky" for x in vs)
